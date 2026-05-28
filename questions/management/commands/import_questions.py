import re
import openpyxl
import docx
from django.core.management.base import BaseCommand
from questions.models import Category, Question, QuestionParameter, Choice


def parse_question_cell(text):
    """Return (orig_number, stem, choices_dict) from a question cell."""
    text = text.strip()

    # Extract question number from start
    m = re.match(r'^(\d+)[\.\,]\s*', text)
    if not m:
        return None, text, {}
    number = int(m.group(1))
    body = text[m.end():]

    # Find where choices block starts: newline then "1."
    choice_start = re.search(r'\n\s*1\.', body)
    if not choice_start:
        return number, body.strip(), {}

    stem = body[:choice_start.start()].strip()
    choice_block = body[choice_start.start():]

    # Split choice block on markers: (newline|tab|2+spaces) followed by digit 1-5 and "."
    marker_re = re.compile(r'(?:\n|\t| {2,})([1-5])\.\s*')
    markers = list(marker_re.finditer(choice_block))

    choices = {}
    for i, marker in enumerate(markers):
        num = int(marker.group(1))
        start = marker.end()
        end = markers[i + 1].start() if i + 1 < len(markers) else len(choice_block)
        choice_body = choice_block[start:end].strip()
        choices[num] = choice_body

    return number, stem, choices


class Command(BaseCommand):
    help = 'Re-import questions from .docx and parameters from .xlsx — match by position order'

    def add_arguments(self, parser):
        parser.add_argument('docx_path', type=str, help='Path to questions .docx')
        parser.add_argument('xlsx_path', type=str, help='Path to parameters .xlsx')
        parser.add_argument('--no-clear', action='store_true', help='Skip clearing existing data')
        parser.add_argument('--num-choices', type=int, default=5, help='Number of choices per question (default: 5)')

    def handle(self, *args, **options):
        if not options['no_clear']:
            self.stdout.write('Clearing existing questions...')
            Choice.objects.all().delete()
            QuestionParameter.objects.all().delete()
            Question.objects.all().delete()
            Category.objects.all().delete()

        # --- Load parameters from Excel by position (row index = question order) ---
        params_by_pos = []  # index 0 = first question
        wb = openpyxl.load_workbook(options['xlsx_path'])
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            params_by_pos.append({
                'param_a': float(row[1]) if row[1] is not None else None,
                'difficulty': float(row[2]) if row[2] is not None else None,
                'param_c': float(row[3]) if row[3] is not None else None,
                'correct_answer': str(int(row[4])) if row[4] is not None else '',
            })
        self.stdout.write(f'Loaded parameters for {len(params_by_pos)} questions from Excel')

        # --- Parse docx in order ---
        doc = docx.Document(options['docx_path'])
        table = doc.tables[0]
        category_cache = {}
        imported = 0
        skipped = 0
        pos = 0  # position counter matching Excel rows

        for row in table.rows[2:]:
            cat_text = row.cells[0].text.strip()
            q_text = row.cells[1].text.strip()
            if not q_text:
                continue

            orig_number, stem, choices = parse_question_cell(q_text)
            if orig_number is None:
                skipped += 1
                continue

            # Use sequential number (pos+1) so it matches Excel row position
            seq_number = pos + 1

            # Category
            if cat_text not in category_cache:
                cat, _ = Category.objects.get_or_create(name=cat_text)
                category_cache[cat_text] = cat

            # Question — use sequential number as the key
            q = Question.objects.create(
                number=seq_number,
                category=category_cache[cat_text],
                body=q_text,
                stem=stem,
            )

            # Choices — ensure all 5 slots exist (image-based choices have empty body)
            num_choices = options['num_choices']
            for choice_num in range(1, num_choices + 1):
                Choice.objects.create(
                    question=q,
                    number=choice_num,
                    body=choices.get(choice_num, ''),
                )

            # Parameters — match by position
            p = params_by_pos[pos] if pos < len(params_by_pos) else {}
            QuestionParameter.objects.create(
                question=q,
                correct_answer=p.get('correct_answer', ''),
                difficulty=p.get('difficulty'),
                param_a=p.get('param_a'),
                param_c=p.get('param_c'),
            )

            pos += 1
            imported += 1

        self.stdout.write(self.style.SUCCESS(
            f'Done: imported {imported} questions, skipped {skipped}'
        ))

        # Summary
        no_answer = QuestionParameter.objects.filter(correct_answer='').count()
        total_choices = Choice.objects.count()
        self.stdout.write(f'  Questions: {imported}, Choices: {total_choices}, No-answer: {no_answer}')
